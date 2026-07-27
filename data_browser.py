"""The data browser: a paginated, sortable table view over a class's view with
typed + faceted search, plus the column-filter, export and import-preview dialogs."""
import os
import re
import csv
import ast
import shutil
import datetime
import sqlite3
import qtawesome as qta
from PySide6.QtWidgets import (
    QWidget, QDialog, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit, QPushButton,
    QComboBox, QStackedWidget, QSizePolicy, QDateTimeEdit, QTableView, QHeaderView,
    QMenu, QMessageBox, QFileDialog, QListWidget, QListWidgetItem, QCheckBox,
    QGroupBox, QRadioButton, QDialogButtonBox, QTextBrowser,
)
from PySide6.QtCore import Qt, QDateTime, QRegularExpression, QUrl
from PySide6.QtSql import QSqlDatabase, QSqlQueryModel, QSqlQuery
from PySide6.QtGui import QRegularExpressionValidator, QDesktopServices

from core import (
    get_app_icon, get_db_path, sync_physical_table, safe_convert, sanitize_name, qid, FILE_PREFIX_LEN,
    db_session,
    files_dir_for, make_stored_filename, display_file_name,
    resolve_file_path, trash_stored_file,
)
from object_editor import ObjectEditorDialog


class ExportOptionsDialog(QDialog):
    """Lets the user pick what to export and in which shape."""
    def __init__(self, parent=None, has_filter=False):
        super().__init__(parent)
        self.setWindowTitle("Export Options")
        self.setWindowIcon(get_app_icon())
        self.resize(440, 360)

        layout = QVBoxLayout(self)

        layout_box = QGroupBox("Layout")
        lb = QVBoxLayout(layout_box)
        self.rb_layout_import = QRadioButton("Import-ready (round-trip)")
        self.rb_layout_import.setChecked(True)
        self.rb_layout_import.setToolTip("Attribute names as headers + relationship IDs. This file can be re-imported.")
        self.rb_layout_full = QRadioButton("Full display view (for reading)")
        self.rb_layout_full.setToolTip("Everything currently visible, including titles and incoming links. Not re-importable.")
        lb.addWidget(self.rb_layout_import)
        lb.addWidget(self.rb_layout_full)
        layout.addWidget(layout_box)

        rows_box = QGroupBox("Rows")
        rb = QVBoxLayout(rows_box)
        self.rb_rows_all = QRadioButton("All rows")
        self.rb_rows_all.setChecked(True)
        self.rb_rows_filter = QRadioButton("Current filter / search only")
        self.rb_rows_filter.setEnabled(has_filter)
        if not has_filter:
            self.rb_rows_filter.setToolTip("No active search filter.")
        rb.addWidget(self.rb_rows_all)
        rb.addWidget(self.rb_rows_filter)
        layout.addWidget(rows_box)

        fmt_box = QGroupBox("Format")
        fb = QHBoxLayout(fmt_box)
        self.rb_fmt_xlsx = QRadioButton("Excel (.xlsx)")
        self.rb_fmt_xlsx.setChecked(True)
        self.rb_fmt_csv = QRadioButton("CSV (.csv)")
        fb.addWidget(self.rb_fmt_xlsx)
        fb.addWidget(self.rb_fmt_csv)
        layout.addWidget(fmt_box)

        self.cb_template = QCheckBox("Template only (column headers, no data)")
        self.cb_template.setToolTip("Produce an empty file with the correct headers to fill in and import.")
        layout.addWidget(self.cb_template)

        buttons = QDialogButtonBox()
        btn_export = buttons.addButton(" Export...", QDialogButtonBox.AcceptRole)
        btn_export.setIcon(qta.icon('fa5s.file-export'))
        buttons.addButton(QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addStretch()
        layout.addWidget(buttons)

    def get_options(self):
        return {
            'import_ready': self.rb_layout_import.isChecked(),
            'rows_filtered': self.rb_rows_filter.isChecked() and self.rb_rows_filter.isEnabled(),
            'fmt': 'xlsx' if self.rb_fmt_xlsx.isChecked() else 'csv',
            'template': self.cb_template.isChecked(),
        }


class ImportPreviewDialog(QDialog):
    """Shows how the chosen file maps onto the class before anything is written."""
    def __init__(self, parent, analysis):
        super().__init__(parent)
        self.setWindowTitle("Import Preview")
        self.setWindowIcon(get_app_icon())
        self.resize(560, 480)

        layout = QVBoxLayout(self)
        report = QTextBrowser()
        report.setOpenExternalLinks(False)

        html = [f"<p><b>File:</b> {analysis['file']}<br><b>Data rows:</b> {analysis['data_count']}</p>"]

        if analysis['problems']:
            html.append("<p style='color:#c0392b;'><b>Cannot import yet:</b></p><ul>")
            for p in analysis['problems']:
                html.append(f"<li style='color:#c0392b;'>{p}</li>")
            html.append("</ul>")

        html.append("<p><b>Columns detected:</b></p><ul>")
        kind_color = {"ok": "#2e7d32", "key": "#1565c0", "ignore": "#888888"}
        for disp, label, kind in analysis['columns']:
            color = kind_color.get(kind, "#000000")
            html.append(f"<li><b>{disp}</b> &rarr; <span style='color:{color};'>{label}</span></li>")
        html.append("</ul>")

        html.append(
            "<p style='color:#555;'>Rows with a value in the <b>ID</b> column update the matching "
            "object; rows with an empty ID are added as new. Columns marked <i>ignored</i> are skipped. "
            "Relationship cells accept comma-separated IDs (or a single-title name).</p>"
        )
        report.setHtml("".join(html))
        layout.addWidget(report)

        buttons = QDialogButtonBox()
        self.btn_import = buttons.addButton(" Import", QDialogButtonBox.AcceptRole)
        self.btn_import.setIcon(qta.icon('fa5s.file-import'))
        self.btn_import.setEnabled(analysis['can_import'])
        buttons.addButton(QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)


class ColumnFilterDialog(QDialog):
    """Excel-style 'filter by values' picker for one column.

    Lists the column's distinct values as a checkable list (plus a Blanks entry),
    with a search box and select-all/clear. Returns the ticked values and whether
    blanks are included; the data browser turns that into an  IN (...)  filter.
    Values start unticked so ticking A and B means 'keep only A or B'.
    """
    MAX_VALUES = 1000

    def __init__(self, parent, db_path, view_name, col_display, kind,
                 options=None, preselected=None, blanks_selected=False):
        super().__init__(parent)
        self.setWindowTitle(f"Filter: {col_display}")
        self.setWindowIcon(get_app_icon())
        self.resize(340, 460)
        self.kind = kind
        self._values = []  # real (typed) value behind each list row

        pre = {str(v) for v in (preselected or [])}

        layout = QVBoxLayout(self)
        layout.addWidget(QLabel(f"<b>{col_display}</b> — tick the values to keep:"))

        self.search = QLineEdit()
        self.search.setPlaceholderText("Find value...")
        self.search.textChanged.connect(self._filter_items)
        layout.addWidget(self.search)

        row = QHBoxLayout()
        btn_all = QPushButton("Select all"); btn_all.clicked.connect(lambda: self._set_all(True))
        btn_none = QPushButton("Clear"); btn_none.clicked.connect(lambda: self._set_all(False))
        row.addWidget(btn_all); row.addWidget(btn_none); row.addStretch()
        layout.addLayout(row)

        self.listw = QListWidget()
        layout.addWidget(self.listw, 1)

        distinct, truncated = self._load_distinct(db_path, view_name, col_display, options)
        if truncated:
            note = QLabel(f"<span style='color:#c0392b;'>Showing the first {self.MAX_VALUES} values; "
                          "narrow with the search box above.</span>")
            note.setWordWrap(True)
            layout.addWidget(note)

        for v in distinct:
            it = QListWidgetItem(self._display(v))
            it.setFlags(it.flags() | Qt.ItemIsUserCheckable)
            it.setCheckState(Qt.Checked if str(v) in pre else Qt.Unchecked)
            it.setData(Qt.UserRole, len(self._values))
            self._values.append(v)
            self.listw.addItem(it)

        self.cb_blanks = QCheckBox("(Blanks / empty)")
        self.cb_blanks.setChecked(blanks_selected)
        layout.addWidget(self.cb_blanks)

        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        layout.addWidget(bb)

    def _display(self, v):
        if self.kind == 'bool':
            return "Yes" if str(v) in ('1', '1.0', 'True', 'true') else "No"
        return str(v)

    def _load_distinct(self, db_path, view_name, col_display, options):
        """Return (values, truncated). Discrete options are merged in even if unused."""
        rows = []
        try:
            with db_session(db_path) as conn:
                cur = conn.cursor()
                cur.execute(
                    f"SELECT DISTINCT {qid(col_display)} FROM {qid(view_name)} "
                    f"WHERE {qid(col_display)} IS NOT NULL AND CAST({qid(col_display)} AS TEXT) <> '' "
                    f"ORDER BY {qid(col_display)} LIMIT {self.MAX_VALUES + 1}")
                rows = [r[0] for r in cur.fetchall()]
        except sqlite3.OperationalError:
            rows = []
        if options:
            present = set(rows)
            rows.extend(o for o in options if o not in present)
            rows = sorted(rows, key=lambda x: str(x))
        truncated = len(rows) > self.MAX_VALUES
        return rows[:self.MAX_VALUES], truncated

    def _filter_items(self, text):
        t = text.lower()
        for i in range(self.listw.count()):
            it = self.listw.item(i)
            it.setHidden(t not in it.text().lower())

    def _set_all(self, checked):
        state = Qt.Checked if checked else Qt.Unchecked
        for i in range(self.listw.count()):
            it = self.listw.item(i)
            if not it.isHidden():  # only affect what the search currently shows
                it.setCheckState(state)

    def get_selection(self):
        vals = [self._values[self.listw.item(i).data(Qt.UserRole)]
                for i in range(self.listw.count())
                if self.listw.item(i).checkState() == Qt.Checked]
        return vals, self.cb_blanks.isChecked()


class DataBrowserPage(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self.current_class_id = None
        self.current_class_name = None
        
        self.current_view_name = None
        self.current_table_name = None
        self.db_path = ""
        self.hidden_columns_memory = {}
        self.file_columns = {}  # display header -> safe column name, for 'file' attributes
        self.column_kinds = {}  # display header -> {'kind': ..., 'options': set?} for typed search
        self.active_filter = None  # the single applied filter descriptor (see _build_search_where)
        
        self.page_size = 100
        self.current_page = 0
        self.total_pages = 0
        self.total_records = 0
        self.current_sort_col = "ID"
        self.current_sort_order = "ASC"
        
        layout = QVBoxLayout(self)
        layout.setContentsMargins(9, 9, 9, 9)

        self.header_label = QLabel("<h2>&nbsp;</h2>")
        layout.addWidget(self.header_label)

        search_layout = QHBoxLayout()
        self.col_combo = QComboBox()
        self.col_combo.addItem("All Columns", -1)
        
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Filter text...")
        self.search_input.returnPressed.connect(self.trigger_search)

        # VS Code-style search toggles (text mode only; they live on the text page).
        self.btn_match_case = QPushButton("Aa")
        self.btn_match_case.setCheckable(True)
        self.btn_match_case.setFixedWidth(34)
        self.btn_match_case.setToolTip("Match Case")
        self.btn_match_case.toggled.connect(self.trigger_search)

        self.btn_whole_word = QPushButton("ab|")
        self.btn_whole_word.setCheckable(True)
        self.btn_whole_word.setFixedWidth(34)
        self.btn_whole_word.setToolTip("Match Whole Word")
        self.btn_whole_word.toggled.connect(self.trigger_search)

        # Operator selector — repopulated per column type (numeric / date / boolean).
        self.op_combo = QComboBox()
        self.op_combo.setFixedWidth(96)
        self.op_combo.currentIndexChanged.connect(self._on_search_op_changed)

        num_rx = QRegularExpressionValidator(QRegularExpression(r'^-?\d*\.?\d*([eE][-+]?\d+)?$'))

        # --- value editors, one per column kind, swapped through a stacked widget ---
        # text: free-text box + toggles
        self.page_text = QWidget()
        pt = QHBoxLayout(self.page_text); pt.setContentsMargins(0, 0, 0, 0)
        pt.addWidget(self.search_input, 1)
        pt.addWidget(self.btn_match_case)
        pt.addWidget(self.btn_whole_word)

        # numeric: value (+ value2 for 'between')
        self.page_num = QWidget()
        pn = QHBoxLayout(self.page_num); pn.setContentsMargins(0, 0, 0, 0)
        self.num_v1 = QLineEdit(); self.num_v1.setPlaceholderText("Number"); self.num_v1.setValidator(num_rx)
        self.num_v1.returnPressed.connect(self.trigger_search)
        self.num_and = QLabel(" and ")
        self.num_v2 = QLineEdit(); self.num_v2.setPlaceholderText("Number"); self.num_v2.setValidator(num_rx)
        self.num_v2.returnPressed.connect(self.trigger_search)
        self.num_v1.setMaximumWidth(140); self.num_v2.setMaximumWidth(140)
        pn.addWidget(self.num_v1); pn.addWidget(self.num_and); pn.addWidget(self.num_v2); pn.addStretch(1)

        # date: datetime (+ datetime2 for 'between')
        self.page_date = QWidget()
        pdl = QHBoxLayout(self.page_date); pdl.setContentsMargins(0, 0, 0, 0)
        self.date_v1 = QDateTimeEdit(); self.date_v1.setCalendarPopup(True)
        self.date_v1.setDisplayFormat("yyyy-MM-dd HH:mm:ss"); self.date_v1.setDateTime(QDateTime.currentDateTime())
        self.date_and = QLabel(" and ")
        self.date_v2 = QDateTimeEdit(); self.date_v2.setCalendarPopup(True)
        self.date_v2.setDisplayFormat("yyyy-MM-dd HH:mm:ss"); self.date_v2.setDateTime(QDateTime.currentDateTime())
        self.date_v1.setMaximumWidth(175); self.date_v2.setMaximumWidth(175)
        pdl.addWidget(self.date_v1); pdl.addWidget(self.date_and); pdl.addWidget(self.date_v2); pdl.addStretch(1)

        # boolean: the operator alone carries the meaning
        self.page_bool = QWidget()
        pb = QHBoxLayout(self.page_bool); pb.setContentsMargins(0, 0, 0, 0)
        pb.addWidget(QLabel("<span style='color:#888;'>choose true / false / empty</span>")); pb.addStretch(1)

        # discrete / set: a button that opens the value picker
        self.page_set = QWidget()
        ps = QHBoxLayout(self.page_set); ps.setContentsMargins(0, 0, 0, 0)
        self.btn_facet = QPushButton(" Choose values…")
        self.btn_facet.setIcon(qta.icon('fa5s.filter'))
        self.btn_facet.clicked.connect(lambda: self.open_facet_dialog())
        ps.addWidget(self.btn_facet); ps.addStretch(1)

        self.value_stack = QStackedWidget()
        for w in (self.page_text, self.page_num, self.page_date, self.page_bool, self.page_set):
            self.value_stack.addWidget(w)
        # Fill width, but stay one row tall — otherwise the stack expands vertically and
        # steals space from the table, leaving the search controls floating in a tall band.
        self.value_stack.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)

        self.col_combo.currentIndexChanged.connect(self._on_search_column_changed)

        self.btn_search = QPushButton(" Search")
        self.btn_search.setIcon(qta.icon('fa5s.search'))
        self.btn_search.clicked.connect(self.trigger_search)

        self.btn_clear_filter = QPushButton()
        self.btn_clear_filter.setIcon(qta.icon('fa5s.times-circle', color='#ff4c4c'))
        self.btn_clear_filter.setToolTip("Clear the active filter")
        self.btn_clear_filter.setFixedWidth(32)
        self.btn_clear_filter.clicked.connect(self.clear_filter)
        self.btn_clear_filter.setVisible(False)
        
        self.btn_import = QPushButton(" Import")
        self.btn_import.setIcon(qta.icon('fa5s.file-import'))
        self.btn_import.clicked.connect(self.import_data)
        
        self.btn_export = QPushButton(" Export")
        self.btn_export.setIcon(qta.icon('fa5s.file-export'))
        self.btn_export.clicked.connect(self.export_data)
        
        self.btn_add = QPushButton(" Add Object")
        self.btn_add.setIcon(qta.icon('fa5s.plus'))
        
        self.btn_edit = QPushButton(" Edit")
        self.btn_edit.setIcon(qta.icon('fa5s.edit'))
        
        self.btn_delete = QPushButton(" Delete")
        self.btn_delete.setIcon(qta.icon('fa5s.trash-alt'))
        
        self.btn_import.setEnabled(False)
        self.btn_add.setEnabled(False) 
        self.btn_edit.setEnabled(False) 
        self.btn_delete.setEnabled(False) 
        
        self.btn_add.clicked.connect(self.open_add_dialog)
        self.btn_edit.clicked.connect(self.open_edit_dialog)
        self.btn_delete.clicked.connect(self.delete_selected)
        
        search_layout.addWidget(QLabel("Search In:"))
        search_layout.addWidget(self.col_combo)
        search_layout.addWidget(self.op_combo)
        search_layout.addWidget(self.value_stack, 1)
        search_layout.addWidget(self.btn_search)
        search_layout.addWidget(self.btn_clear_filter)
        layout.addLayout(search_layout)

        # Data actions live on their own row so the filter row always has room.
        action_layout = QHBoxLayout()
        for _b in (self.btn_import, self.btn_export, self.btn_add, self.btn_edit, self.btn_delete):
            action_layout.addWidget(_b)
        action_layout.addStretch(1)
        layout.addLayout(action_layout)

        # Start in 'All Columns' free-text mode.
        self._configure_filter_ui('all')

        self.table_view = QTableView()
        self.table_view.setAlternatingRowColors(True)
        self.table_view.setSelectionBehavior(QTableView.SelectRows)
        self.table_view.setSelectionMode(QTableView.SingleSelection)
        self.table_view.setEditTriggers(QTableView.NoEditTriggers)
        self.table_view.doubleClicked.connect(self.on_cell_double_clicked)
        self.table_view.setContextMenuPolicy(Qt.CustomContextMenu)
        self.table_view.customContextMenuRequested.connect(self.show_context_menu)
        self.table_view.setSortingEnabled(True)
        
        header = self.table_view.horizontalHeader()
        header.setSectionsMovable(True)
        header.setSectionResizeMode(QHeaderView.Interactive)
        header.setStretchLastSection(True)
        header.sortIndicatorChanged.connect(self.on_sort_changed)
        header.setContextMenuPolicy(Qt.CustomContextMenu)
        header.customContextMenuRequested.connect(self.show_header_context_menu)
        
        layout.addWidget(self.table_view)

        self.query_model = QSqlQueryModel()
        self.table_view.setModel(self.query_model)
        self.table_view.selectionModel().selectionChanged.connect(self.on_selection_changed)

        page_layout = QHBoxLayout()
        
        self.btn_prev = QPushButton(" Prev") 
        self.btn_prev.setIcon(qta.icon('fa5s.chevron-left'))
        self.btn_prev.clicked.connect(self.prev_page)
        
        self.page_label = QLabel("Page 0 of 0 (Total: 0)")
        self.page_label.setAlignment(Qt.AlignCenter)
        
        self.btn_next = QPushButton("Next ") 
        self.btn_next.setIcon(qta.icon('fa5s.chevron-right'))
        self.btn_next.setLayoutDirection(Qt.RightToLeft)
        self.btn_next.clicked.connect(self.next_page)
        
        self.page_size_combo = QComboBox()
        self.page_size_combo.addItems(["50", "100", "500", "1000"])
        self.page_size_combo.setCurrentText("100")
        self.page_size_combo.currentTextChanged.connect(self.change_page_size)
        
        page_layout.addStretch()
        page_layout.addWidget(self.btn_prev)
        page_layout.addWidget(self.page_label)
        page_layout.addWidget(self.btn_next)
        page_layout.addWidget(QLabel("  Rows per page:"))
        page_layout.addWidget(self.page_size_combo)
        layout.addLayout(page_layout)

    def clear_view(self):
        """Blank the data view: clear the grid AND the class-name heading together.

        A schema sync empties the query model so it doesn't hold a read lock during
        DDL; resetting the heading here keeps the two in lock-step so the bold class
        name can't linger above an empty table.
        """
        self.query_model.clear()
        self.header_label.setText("<h2>&nbsp;</h2>")

    def _build_column_kinds(self, cur, class_id, columns):
        """Classify each view column so the search bar can offer typed operators.

        Real attribute columns get their type; the ID column is numeric; relationship
        and other computed columns are 'multivalue' (concatenated -> text search only).
        """
        cur.execute("SELECT name, data_type, lookup_query FROM attributes WHERE class_id = ?", (class_id,))
        attr_meta = {name: (dtype, lookup) for name, dtype, lookup in cur.fetchall()}
        kinds = {}
        for col in columns:
            if col == "ID":
                kinds[col] = {'kind': 'num'}
                continue
            meta = attr_meta.get(col)
            if not meta:
                kinds[col] = {'kind': 'multivalue'}  # relationship / computed column
                continue
            dtype, lookup = meta
            if dtype in ('int', 'float'):
                kinds[col] = {'kind': 'num'}
            elif dtype == 'date':
                kinds[col] = {'kind': 'date'}
            elif dtype == 'boolean':
                kinds[col] = {'kind': 'bool'}
            elif dtype == 'discrete':
                try:
                    opts = {r[0] for r in cur.execute(
                        "SELECT value FROM discrete_options WHERE type_id = ? ORDER BY row_order",
                        (int(lookup),)).fetchall()}
                except (ValueError, TypeError, sqlite3.OperationalError):
                    opts = set()
                kinds[col] = {'kind': 'discrete', 'options': opts}
            else:  # string, long string, list, matrix, file, look-through
                kinds[col] = {'kind': 'text'}
        return kinds

    def load_table_data(self, class_id, class_name):
        self.current_class_id = class_id
        self.current_class_name = class_name
        self.header_label.setText(f"<h2>{class_name}</h2>")
        
        self.btn_import.setEnabled(True)
        self.btn_add.setEnabled(True)
        self.btn_edit.setEnabled(False)
        self.btn_delete.setEnabled(False)
        
        self.db_path = get_db_path()

        # Release any read lock the Qt view holds on the old class's view before the
        # schema sync runs DDL on this database file (belt-and-suspenders alongside WAL).
        self.query_model.clear()

        try:
            self.current_view_name, self.current_table_name = sync_physical_table(self.db_path, class_id, class_name, parent_widget=self)
            
            if not self.current_view_name:
                raise RuntimeError("Sync aborted or failed to generate a valid view.")
                
            with db_session(self.db_path) as conn:
                cur = conn.cursor()
                cur.execute(f"PRAGMA table_info({qid(self.current_view_name)})")
                columns = [row[1] for row in cur.fetchall()]

                cur.execute("SELECT name FROM attributes WHERE class_id = ? AND data_type = 'file'", (class_id,))
                self.file_columns = {row[0]: sanitize_name(row[0]) for row in cur.fetchall()}

                self.column_kinds = self._build_column_kinds(cur, class_id, columns)

            if not columns:
                raise sqlite3.OperationalError(f"View '{self.current_view_name}' is broken and returned no columns.")
                
        except Exception as e:
            QMessageBox.critical(self, "Schema Load Error", f"A database operational error occurred while loading '{class_name}'.\n\nThis is usually caused by broken 'Look-through' attributes referencing deleted or missing columns.\n\nDetails:\n{str(e)}")
            self.header_label.setText("<h2>Sync Aborted</h2>")
            self.btn_import.setEnabled(False)
            self.btn_add.setEnabled(False)
            return
            
        self._reset_filter_ui()
        self.current_sort_col = "ID"
        self.current_sort_order = "ASC"
        self.current_page = 0
        
        header = self.table_view.horizontalHeader()
        header.blockSignals(True)
        header.setSortIndicator(0, Qt.AscendingOrder)
        header.blockSignals(False)

        self.col_combo.blockSignals(True)
        self.col_combo.clear()
        self.col_combo.addItem("All Columns", -1)
        for i, col_name in enumerate(columns):
            self.col_combo.addItem(col_name, i)
        self.col_combo.blockSignals(False)
        self._configure_filter_ui('all')

        self.build_and_exec_query()

    @staticmethod
    def _glob_escape(s):
        # GLOB has no ESCAPE clause; wrap each metachar in a one-char class to match it literally.
        out = []
        for ch in s:
            out.append(f"[{ch}]" if ch in "*?[" else ch)
        return "".join(out)

    @staticmethod
    def _like_escape(s):
        return s.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")

    def _build_text_condition(self, col_sql, search_text, match_case, whole_word):
        """(sql, params) for a 'contains' text match honouring Match Case / Whole Word.

        Used by both the COUNT and the data query so paging stays consistent. Reads
        no widgets — case/word come in as args — so it's plain SQLite SQL that works
        on either connection. Case-insensitive matching is ASCII-only (as LIKE was).
        """
        col_expr = f"CAST({col_sql} AS TEXT)"
        if whole_word:
            boundary = "[^A-Za-z0-9_]"
            term = self._glob_escape(search_text)
            if not match_case:
                col_expr = f"LOWER({col_expr})"
                term = term.lower()
            patterns = [
                term,                                      # whole value equals the word
                f"{term}{boundary}*",                      # word at the start
                f"*{boundary}{term}",                      # word at the end
                f"*{boundary}{term}{boundary}*",           # word in the middle
            ]
            cond = "(" + " OR ".join([f"{col_expr} GLOB ?"] * len(patterns)) + ")"
            return cond, patterns
        if match_case:
            # INSTR is case- and accent-sensitive (binary), i.e. true "match case".
            return f"INSTR({col_expr}, ?) > 0", [search_text]
        esc = self._like_escape(search_text)
        return f"{col_expr} LIKE ? ESCAPE '\\'", [f"%{esc}%"]

    def _num_condition(self, col, op, v, v2=None):
        ce = qid(col)
        ops = {'=': '=', '!=': '<>', '>': '>', '>=': '>=', '<': '<', '<=': '<='}
        if op == 'empty':
            return f"{ce} IS NULL", []
        if op == 'between':
            if v is None or v2 is None:
                return "", []
            lo, hi = (v, v2) if v <= v2 else (v2, v)
            return f"{ce} BETWEEN ? AND ?", [lo, hi]
        if v is None or op not in ops:
            return "", []
        return f"{ce} {ops[op]} ?", [v]

    def _date_condition(self, col, op, v, v2=None):
        # Dates are stored zero-padded 'yyyy-MM-dd HH:mm:ss', so string order == time order.
        ce = qid(col)
        if op == 'empty':
            return f"({ce} IS NULL OR CAST({ce} AS TEXT) = '')", []
        if op == 'on':
            return f"date({ce}) = date(?)", [v]
        if op == 'before':
            return f"{ce} < ?", [v]
        if op == 'after':
            return f"{ce} > ?", [v]
        if op == 'between':
            if v is None or v2 is None:
                return "", []
            lo, hi = (v, v2) if v <= v2 else (v2, v)
            return f"({ce} >= ? AND {ce} <= ?)", [lo, hi]
        return "", []

    def _bool_condition(self, col, op):
        ce = qid(col)
        if op == 'true':
            return f"{ce} = 1", []
        if op == 'false':
            return f"{ce} = 0", []
        if op == 'empty':
            return f"{ce} IS NULL", []
        return "", []

    def _set_condition(self, col, values, blanks):
        ce = qid(col)
        parts, params = [], []
        if values:
            ph = ", ".join(["?"] * len(values))
            parts.append(f"{ce} IN ({ph})")
            params.extend(values)
        if blanks:
            parts.append(f"({ce} IS NULL OR CAST({ce} AS TEXT) = '')")
        if not parts:
            return "", []
        return "(" + " OR ".join(parts) + ")", params

    def _build_search_where(self, cur):
        """Turn the single active filter descriptor into (where_sql, params).

        where_sql is '' or starts with ' WHERE '. Reads only self.active_filter,
        never the widgets, so the COUNT query, the data query and 'current filter
        only' exports all agree — and the logic stays unit-testable without a UI.
        """
        f = self.active_filter
        if not f:
            return "", []
        kind = f.get('kind')
        cond, params = "", []

        if kind == 'text':
            text = (f.get('text') or "").strip()
            if not text:
                return "", []
            mc, ww = f.get('match_case', False), f.get('whole_word', False)
            if f.get('scope') == 'all':
                cur.execute(f"PRAGMA table_info({qid(self.current_view_name)})")
                names = [row[1] for row in cur.fetchall()]
                ors = []
                for c in names:
                    cc, cp = self._build_text_condition(qid(c), text, mc, ww)
                    ors.append(cc)
                    params.extend(cp)
                if not ors:
                    return "", []
                cond = "(" + " OR ".join(ors) + ")"
            else:
                cond, params = self._build_text_condition(qid(f['col']), text, mc, ww)
        elif kind == 'num':
            cond, params = self._num_condition(f['col'], f['op'], f.get('v'), f.get('v2'))
        elif kind == 'date':
            cond, params = self._date_condition(f['col'], f['op'], f.get('v'), f.get('v2'))
        elif kind == 'bool':
            cond, params = self._bool_condition(f['col'], f['op'])
        elif kind == 'set':
            cond, params = self._set_condition(f['col'], f.get('values', []), f.get('blanks', False))

        if not cond:
            return "", []
        return " WHERE " + cond, params

    # ---------------- search bar wiring ----------------
    def _kind_of(self, col_display):
        info = self.column_kinds.get(col_display)
        return info['kind'] if info else 'text'

    @staticmethod
    def _parse_float(s):
        s = (s or "").strip()
        if s == "":
            return None
        try:
            return float(s)
        except ValueError:
            return None

    def _on_search_column_changed(self, *_):
        data = self.col_combo.currentData()
        kind = 'all' if (data is None or data == -1) else self._kind_of(self.col_combo.currentText())
        self._configure_filter_ui(kind)

    def _configure_filter_ui(self, kind):
        """Repopulate the operator combo and swap the value editor for this column kind."""
        self.op_combo.blockSignals(True)
        self.op_combo.clear()
        if kind == 'num':
            for lbl, code in [("=", '='), ("≠", '!='), (">", '>'), ("≥", '>='),
                              ("<", '<'), ("≤", '<='), ("between", 'between'), ("is empty", 'empty')]:
                self.op_combo.addItem(lbl, code)
            self.op_combo.setVisible(True)
            self.value_stack.setCurrentWidget(self.page_num)
        elif kind == 'date':
            for lbl, code in [("on", 'on'), ("before", 'before'), ("after", 'after'),
                              ("between", 'between'), ("is empty", 'empty')]:
                self.op_combo.addItem(lbl, code)
            self.op_combo.setVisible(True)
            self.value_stack.setCurrentWidget(self.page_date)
        elif kind == 'bool':
            for lbl, code in [("is true", 'true'), ("is false", 'false'), ("is empty", 'empty')]:
                self.op_combo.addItem(lbl, code)
            self.op_combo.setVisible(True)
            self.value_stack.setCurrentWidget(self.page_bool)
        elif kind == 'discrete':
            self.op_combo.setVisible(False)
            self.value_stack.setCurrentWidget(self.page_set)
        else:  # all / text / multivalue / other
            self.op_combo.setVisible(False)
            self.value_stack.setCurrentWidget(self.page_text)
        self.op_combo.blockSignals(False)
        self._on_search_op_changed()

    def _on_search_op_changed(self, *_):
        op = self.op_combo.currentData()
        page = self.value_stack.currentWidget()
        if page is self.page_num:
            self.num_v1.setVisible(op != 'empty')
            self.num_and.setVisible(op == 'between')
            self.num_v2.setVisible(op == 'between')
        elif page is self.page_date:
            self.date_v1.setVisible(op != 'empty')
            self.date_and.setVisible(op == 'between')
            self.date_v2.setVisible(op == 'between')

    def _compose_filter_from_ui(self):
        """Build a filter descriptor from the bar widgets.

        Returns a descriptor, None (clear the filter), or "KEEP" (leave the current
        filter untouched — used for discrete columns, whose filter is set via the
        value picker rather than the Search button).
        """
        data = self.col_combo.currentData()
        if data is None or data == -1:
            text = self.search_input.text().strip()
            return {'kind': 'text', 'scope': 'all', 'text': text,
                    'match_case': self.btn_match_case.isChecked(),
                    'whole_word': self.btn_whole_word.isChecked()} if text else None

        col = self.col_combo.currentText()
        kind = self._kind_of(col)

        if kind == 'num':
            op = self.op_combo.currentData()
            if op == 'empty':
                return {'kind': 'num', 'col': col, 'op': 'empty'}
            v = self._parse_float(self.num_v1.text())
            if op == 'between':
                v2 = self._parse_float(self.num_v2.text())
                return {'kind': 'num', 'col': col, 'op': 'between', 'v': v, 'v2': v2} if (v is not None and v2 is not None) else None
            return {'kind': 'num', 'col': col, 'op': op, 'v': v} if v is not None else None

        if kind == 'date':
            op = self.op_combo.currentData()
            if op == 'empty':
                return {'kind': 'date', 'col': col, 'op': 'empty'}
            v = self.date_v1.dateTime().toString("yyyy-MM-dd HH:mm:ss")
            if op == 'between':
                v2 = self.date_v2.dateTime().toString("yyyy-MM-dd HH:mm:ss")
                return {'kind': 'date', 'col': col, 'op': 'between', 'v': v, 'v2': v2}
            return {'kind': 'date', 'col': col, 'op': op, 'v': v}

        if kind == 'bool':
            return {'kind': 'bool', 'col': col, 'op': self.op_combo.currentData()}

        if kind == 'discrete':
            return "KEEP"  # driven by the value picker, not the Search button

        # text / multivalue / other -> 'contains' on this one column
        text = self.search_input.text().strip()
        return {'kind': 'text', 'scope': 'col', 'col': col, 'text': text,
                'match_case': self.btn_match_case.isChecked(),
                'whole_word': self.btn_whole_word.isChecked()} if text else None

    def open_facet_dialog(self, col=None):
        if not self.current_view_name:
            return
        if not col:
            col = self.col_combo.currentText()
        info = self.column_kinds.get(col, {'kind': 'text'})
        kind = info['kind']
        if kind in ('multivalue', 'other'):
            QMessageBox.information(self, "Filter by values",
                f"'{col}' combines several values per cell (a relationship/computed column), "
                "so an exact value filter doesn't apply here.\n\nUse the text search instead.")
            return
        pre, blanks = [], False
        if self.active_filter and self.active_filter.get('kind') == 'set' and self.active_filter.get('col') == col:
            pre = self.active_filter.get('values', [])
            blanks = self.active_filter.get('blanks', False)
        dlg = ColumnFilterDialog(self, self.db_path, self.current_view_name, col, kind,
                                 info.get('options'), pre, blanks)
        if not dlg.exec():
            return
        values, blanks_sel = dlg.get_selection()
        self.active_filter = ({'kind': 'set', 'col': col, 'values': values, 'blanks': blanks_sel}
                              if (values or blanks_sel) else None)
        idx = self.col_combo.findText(col)
        if idx >= 0:
            self.col_combo.setCurrentIndex(idx)
        self.current_page = 0
        self.build_and_exec_query()
        self._update_clear_btn()

    def clear_filter(self):
        self.active_filter = None
        self.search_input.clear()
        self.num_v1.clear()
        self.num_v2.clear()
        self.current_page = 0
        self.build_and_exec_query()
        self._update_clear_btn()

    def _update_clear_btn(self):
        self.btn_clear_filter.setVisible(self.active_filter is not None)

    def _reset_filter_ui(self):
        self.active_filter = None
        self.search_input.clear()
        self.num_v1.clear()
        self.num_v2.clear()
        self.btn_clear_filter.setVisible(False)

    def _get_import_schema(self, cur):
        """The canonical, round-trippable column schema for this class.

        Returns (attributes, relationships, lookthrough_names) where the import-ready
        header order is ['ID'] + [a.name for attributes] + [r.name for relationships].
        """
        cur.execute("SELECT id, name, data_type, is_unique, is_required, lookup_query FROM attributes WHERE class_id = ? ORDER BY row_order", (self.current_class_id,))
        attr_rows = cur.fetchall()

        cur.execute("""
            SELECT a.name, COUNT(m.id) FROM attributes a
            JOIN matrix_columns m ON a.id = m.attribute_id
            WHERE a.class_id = ? GROUP BY a.id
        """, (self.current_class_id,))
        matrix_counts = {r[0]: r[1] for r in cur.fetchall()}

        attributes = []
        lookthrough_names = []
        for aid, name, dtype, uniq, req, lookup_query in attr_rows:
            if dtype == "look-through":
                lookthrough_names.append(name)
                continue
            options = None
            if dtype == "discrete":
                try:
                    options = {r[0] for r in cur.execute(
                        "SELECT value FROM discrete_options WHERE type_id = ?", (int(lookup_query),)).fetchall()}
                except (ValueError, TypeError, sqlite3.OperationalError):
                    options = set()
            attributes.append({
                "name": name, "safe": sanitize_name(name), "type": dtype,
                "unique": bool(uniq), "required": bool(req),
                "matrix_count": matrix_counts.get(name, 0),
                "options": options,
            })

        cur.execute("""
            SELECT c.id, c.name, r.rel_type, r.is_required FROM relationships r
            JOIN classes c ON r.target_class = c.id
            WHERE r.source_class = ? ORDER BY r.row_order
        """, (self.current_class_id,))
        relationships = []
        seen = set()
        for tcid, tname, rel_type, rel_required in cur.fetchall():
            safe_t = sanitize_name(tname)
            if safe_t in seen:
                continue
            seen.add(safe_t)
            cur.execute("SELECT name FROM attributes WHERE class_id = ? AND is_title = 1 ORDER BY row_order", (tcid,))
            title_cols = [r[0] for r in cur.fetchall()]
            relationships.append({
                "name": tname, "safe": safe_t, "rel_type": rel_type,
                "required": bool(rel_required),
                "target_table": f"objects_{safe_t}",
                "junc_table": f"rel_{self.current_table_name}_to_objects_{safe_t}",
                "base_view": f"base_view_objects_{safe_t}",
                "title_cols": title_cols,
            })
        return attributes, relationships, lookthrough_names

    @staticmethod
    def _cell_to_value(cell):
        """Normalise a spreadsheet/CSV cell to a trimmed string or None."""
        if isinstance(cell, datetime.datetime):
            return cell.strftime("%Y-%m-%d %H:%M:%S")
        if cell is None:
            return None
        s = str(cell).strip()
        return s if s != "" else None

    def build_and_exec_query(self):
        if not self.current_view_name: return
        
        with db_session(self.db_path) as conn:
            cur = conn.cursor()

            base_query = f"FROM {qid(self.current_view_name)}"
            where_sql, params = self._build_search_where(cur)

            count_query = f"SELECT COUNT(*) {base_query} {where_sql}"
            cur.execute(count_query, params)
            self.total_records = cur.fetchone()[0]
            self.total_pages = max(1, (self.total_records + self.page_size - 1) // self.page_size)
            
            if self.current_page >= self.total_pages:
                self.current_page = max(0, self.total_pages - 1)
                
            offset = self.current_page * self.page_size
            order_sql = f"ORDER BY {qid(self.current_sort_col)} {self.current_sort_order}"
            final_query = f"SELECT * {base_query} {where_sql} {order_sql} LIMIT {self.page_size} OFFSET {offset}"
            
            db = QSqlDatabase.database()
            query = QSqlQuery(db)
            query.prepare(final_query)
            for p in params: query.addBindValue(p)
            
            query.exec()
            if query.lastError().isValid():
                QMessageBox.critical(self, "Database Query Error", f"An error occurred while fetching data from the view:\n\n{query.lastError().text()}")
                
            self.query_model.setQuery(query)

            # Qt fetches a query lazily (256 rows at a time) and keeps the statement -
            # and with it a read lock - open until the last row is fetched. Under the
            # rollback journal a network database uses, that read lock blocks every
            # other user from saving for as long as this grid sits partially scrolled.
            # The page is capped at page_size rows, so draining it now is bounded work.
            while self.query_model.canFetchMore():
                self.query_model.fetchMore()
            
            hidden_for_class = self.hidden_columns_memory.get(self.current_class_name, set())
            for i in range(self.query_model.columnCount()):
                col_name = self.query_model.headerData(i, Qt.Horizontal)
                if col_name in hidden_for_class:
                    self.table_view.setColumnHidden(i, True)
                else:
                    self.table_view.setColumnHidden(i, False)
            
            self.page_label.setText(f"Page {self.current_page + 1} of {self.total_pages} (Total: {self.total_records})")
            self.btn_prev.setEnabled(self.current_page > 0)
            self.btn_next.setEnabled(self.current_page < self.total_pages - 1)

    def on_sort_changed(self, logical_index, order):
        col_name = self.query_model.headerData(logical_index, Qt.Horizontal)
        if not col_name: return 
        self.current_sort_col = col_name
        self.current_sort_order = "ASC" if order == Qt.AscendingOrder else "DESC"
        self.build_and_exec_query()

    def trigger_search(self):
        composed = self._compose_filter_from_ui()
        if composed != "KEEP":
            self.active_filter = composed
        self.current_page = 0
        self.build_and_exec_query()
        self._update_clear_btn()

    def prev_page(self):
        if self.current_page > 0:
            self.current_page -= 1
            self.build_and_exec_query()
            
    def next_page(self):
        if self.current_page < self.total_pages - 1:
            self.current_page += 1
            self.build_and_exec_query()
            
    def change_page_size(self, new_size):
        self.page_size = int(new_size)
        self.current_page = 0
        self.build_and_exec_query()

    def on_selection_changed(self):
        has_sel = bool(self.table_view.selectionModel().selectedRows())
        self.btn_edit.setEnabled(has_sel)
        self.btn_delete.setEnabled(has_sel)

    def get_selected_id(self):
        rows = self.table_view.selectionModel().selectedRows()
        if not rows: return None
        idx = rows[0]
        id_index = self.query_model.index(idx.row(), 0)
        return self.query_model.data(id_index)

    def show_header_context_menu(self, position):
        header = self.table_view.horizontalHeader()
        logical_index = header.logicalIndexAt(position)
        menu = QMenu()

        col_name = self.query_model.headerData(logical_index, Qt.Horizontal) if logical_index >= 0 else None

        act_filter = None
        if col_name and self.column_kinds.get(col_name, {}).get('kind') not in ('multivalue', 'other'):
            act_filter = menu.addAction(qta.icon('fa5s.filter'), f"Filter by values in '{col_name}'…")

        act_hide = None
        if col_name:
            visible_count = sum(1 for i in range(header.count()) if not self.table_view.isColumnHidden(i))
            if visible_count > 1:
                if act_filter: menu.addSeparator()
                act_hide = menu.addAction(f"Hide '{col_name}' Temporarily")

        hidden_cols = [i for i in range(header.count()) if self.table_view.isColumnHidden(i)]
        act_unhide = None
        if hidden_cols:
            if act_hide or act_filter: menu.addSeparator()
            act_unhide = menu.addAction("Unhide All Columns")

        act_clear = None
        if self.active_filter is not None:
            menu.addSeparator()
            act_clear = menu.addAction("Clear active filter")

        if not menu.actions(): return

        action = menu.exec(header.viewport().mapToGlobal(position))
        if action is None:
            return

        if act_filter and action == act_filter:
            self.open_facet_dialog(col_name)
        elif act_hide and action == act_hide:
            self.table_view.setColumnHidden(logical_index, True)
            if col_name: self.hidden_columns_memory.setdefault(self.current_class_name, set()).add(col_name)
        elif act_unhide and action == act_unhide:
            for i in range(header.count()): self.table_view.setColumnHidden(i, False)
            if self.current_class_name in self.hidden_columns_memory:
                self.hidden_columns_memory[self.current_class_name].clear()
        elif act_clear and action == act_clear:
            self.clear_filter()

    def show_context_menu(self, position):
        if not self.table_view.selectionModel().selectedRows(): return
        from PySide6.QtWidgets import QMenu
        menu = QMenu()
        act_edit = menu.addAction(qta.icon('fa5s.edit'), "Edit Object")
        act_del = menu.addAction(qta.icon('fa5s.trash-alt'), "Delete Object")
        
        action = menu.exec(self.table_view.viewport().mapToGlobal(position))
        if action == act_edit: self.open_edit_dialog()
        elif action == act_del: self.delete_selected()

    def open_add_dialog(self):
        dialog = ObjectEditorDialog(self.db_path, self.current_class_id, self.current_class_name, self.current_table_name, None, self)
        if dialog.exec(): self.build_and_exec_query() 
            
    def on_cell_double_clicked(self, index):
        # Double-clicking a 'file' cell opens/saves the attachment; any other cell edits the row.
        header = self.query_model.headerData(index.column(), Qt.Horizontal)
        if header in self.file_columns:
            self._open_file_cell(index, self.file_columns[header])
        else:
            self.open_edit_dialog()

    def _open_file_cell(self, index, safe_col):
        id_index = self.query_model.index(index.row(), 0)
        obj_id = self.query_model.data(id_index)
        if not obj_id:
            return
        with db_session(self.db_path) as conn:
            row = conn.execute(f"SELECT {qid(safe_col)} FROM {qid(self.current_table_name)} WHERE id = ?", (obj_id,)).fetchone()
        stored = row[0] if row else None
        if not stored:
            QMessageBox.information(self, "No File", "There is no file attached to this cell.")
            return

        path = resolve_file_path(self.db_path, stored)
        display = display_file_name(stored)
        if not path or not os.path.exists(path):
            QMessageBox.warning(self, "File Not Found",
                f"'{display}' is referenced by this record but is no longer in the files folder.")
            return

        box = QMessageBox(self)
        box.setWindowTitle("File")
        box.setIcon(QMessageBox.Question)
        box.setText(f"<b>{display}</b>")
        box.setInformativeText("What would you like to do?")
        btn_open = box.addButton(" Open", QMessageBox.AcceptRole)
        btn_open.setIcon(qta.icon('fa5s.external-link-alt'))
        btn_save = box.addButton(" Save As...", QMessageBox.ActionRole)
        btn_save.setIcon(qta.icon('fa5s.download'))
        box.addButton(QMessageBox.Cancel)
        box.exec()

        clicked = box.clickedButton()
        if clicked == btn_open:
            QDesktopServices.openUrl(QUrl.fromLocalFile(os.path.abspath(path)))
        elif clicked == btn_save:
            dest, _ = QFileDialog.getSaveFileName(self, "Save File As", display, "All Files (*.*)")
            if dest:
                try:
                    shutil.copy2(path, dest)
                except OSError as e:
                    QMessageBox.critical(self, "Save Failed", str(e))

    def open_edit_dialog(self):
        obj_id = self.get_selected_id()
        if not obj_id: return
        dialog = ObjectEditorDialog(self.db_path, self.current_class_id, self.current_class_name, self.current_table_name, obj_id, self)
        if dialog.exec(): self.build_and_exec_query()

    def delete_selected(self):
        obj_id = self.get_selected_id()
        if not obj_id: return
        reply = QMessageBox.question(self, "Confirm Delete", f"Are you sure you want to delete this {self.current_class_name}?\n\nIts relationships will be automatically safely removed.", QMessageBox.Yes | QMessageBox.No)
        if reply == QMessageBox.Yes:
            try:
                # Collect attached files first so we can retire them after a successful delete.
                stored_files = []
                if self.file_columns:
                    with db_session(self.db_path) as conn:
                        cols = list(self.file_columns.values())
                        sel = ", ".join(qid(c) for c in cols)
                        r = conn.execute(f"SELECT {sel} FROM {qid(self.current_table_name)} WHERE id = ?", (obj_id,)).fetchone()
                        if r:
                            stored_files = [v for v in r if v]

                with db_session(self.db_path) as conn:
                    conn.execute("PRAGMA foreign_keys = 1")
                    conn.execute("BEGIN IMMEDIATE")
                    conn.execute(f"DELETE FROM {qid(self.current_table_name)} WHERE id = ?", (obj_id,))
                    conn.commit()

                for stored in stored_files:
                    trash_stored_file(self.db_path, stored)

                self.build_and_exec_query()
            except Exception as e:
                QMessageBox.critical(self, "Error", str(e))

    def export_data(self):
        if not self.current_view_name: return

        has_filter = self.active_filter is not None
        opt_dlg = ExportOptionsDialog(self, has_filter=has_filter)
        if not opt_dlg.exec(): return
        opts = opt_dlg.get_options()

        ext = ".xlsx" if opts['fmt'] == 'xlsx' else ".csv"
        filt = "Excel Files (*.xlsx)" if opts['fmt'] == 'xlsx' else "CSV Files (*.csv)"
        default_name = f"{sanitize_name(self.current_class_name)}{'_template' if opts['template'] else ''}{ext}"
        file_path, _ = QFileDialog.getSaveFileName(self, "Export Data", default_name, filt)
        if not file_path: return
        if not file_path.lower().endswith(ext): file_path += ext

        try:
            headers, rows = self._gather_export_rows(opts)
        except Exception as e:
            QMessageBox.critical(self, "Export Error", f"Could not gather data to export:\n\n{e}")
            return

        try:
            if opts['fmt'] == 'xlsx':
                try:
                    import openpyxl
                except ImportError:
                    file_path = file_path[:-5] + ".csv" if file_path.lower().endswith(".xlsx") else file_path + ".csv"
                    self._write_csv(file_path, headers, rows)
                    QMessageBox.warning(self, "Exported as CSV",
                        f"'openpyxl' is not installed, so the data was exported as CSV instead:\n{file_path}\n\n"
                        "Install it (pip install openpyxl) for Excel export.")
                    return
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.append(headers)
                for row in rows: ws.append(row)
                wb.save(file_path)
            else:
                self._write_csv(file_path, headers, rows)
        except Exception as e:
            QMessageBox.critical(self, "Export Error", f"Failed to write file:\n\n{e}")
            return

        what = "template" if opts['template'] else f"{len(rows)} row(s)"
        QMessageBox.information(self, "Export Complete", f"Exported {what} to:\n{file_path}")

    @staticmethod
    def _write_csv(file_path, headers, rows):
        with open(file_path, mode='w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            writer.writerows(rows)

    def _gather_export_rows(self, opts):
        """Return (headers, rows) for the chosen export options."""
        with db_session(self.db_path) as conn:
            cur = conn.cursor()

            if opts['import_ready']:
                attrs, rels, _lookthroughs = self._get_import_schema(cur)
                headers = ["ID"] + [a['name'] for a in attrs] + [r['name'] for r in rels]
                if opts['template']:
                    return headers, []

                selects = ["m.id"]
                # 'file' columns export the display name (not the raw "<uuid>__name" reference).
                selects += [
                    (f"substr(m.{qid(a['safe'])}, {FILE_PREFIX_LEN + 1})" if a['type'] == 'file' else f"m.{qid(a['safe'])}")
                    for a in attrs
                ]
                selects += [f"(SELECT GROUP_CONCAT(target_id) FROM {qid(r['junc_table'])} WHERE source_id = m.id)" for r in rels]
                sql = f"SELECT {', '.join(selects)} FROM {qid(self.current_table_name)} m"

                params = []
                if opts['rows_filtered']:
                    where_sql, wparams = self._build_search_where(cur)
                    cur.execute(f"SELECT [ID] FROM {qid(self.current_view_name)}{where_sql}", wparams)
                    ids = [r[0] for r in cur.fetchall()]
                    if not ids:
                        return headers, []
                    # The IDs go through a TEMP table rather than one bound parameter
                    # each: SQLite caps parameters per statement (SQLITE_LIMIT_VARIABLE_NUMBER,
                    # 32766 in the bundled build), so "IN (?,?,...)" fails with
                    # 'too many SQL variables' once a filter matches more rows than that.
                    # TEMP tables live in this connection's private temp store, not the
                    # shared database file, so this takes no write lock for other users.
                    cur.execute("CREATE TEMP TABLE IF NOT EXISTS export_ids (id INTEGER PRIMARY KEY)")
                    cur.execute("DELETE FROM export_ids")
                    cur.executemany("INSERT INTO export_ids (id) VALUES (?)", ((i,) for i in ids))
                    sql += " WHERE m.id IN (SELECT id FROM export_ids)"
                sql += " ORDER BY m.id ASC"
                cur.execute(sql, params)
                raw = cur.fetchall()
            else:
                cur.execute(f"PRAGMA table_info({qid(self.current_view_name)})")
                headers = [row[1] for row in cur.fetchall()]
                if opts['template']:
                    return headers, []
                where_sql, params = self._build_search_where(cur) if opts['rows_filtered'] else ("", [])
                order_sql = f"ORDER BY {qid(self.current_sort_col)} {self.current_sort_order}"
                cur.execute(f"SELECT * FROM {qid(self.current_view_name)}{where_sql} {order_sql}", params)
                raw = cur.fetchall()

        return headers, [["" if v is None else str(v) for v in row] for row in raw]

    def _read_table_file(self, file_path):
        """Read an .xlsx or .csv file into a list of row-lists (header row first)."""
        if file_path.lower().endswith(".csv"):
            with open(file_path, newline='', encoding='utf-8-sig') as f:
                return [list(r) for r in csv.reader(f)]
        try:
            import openpyxl
        except ImportError:
            raise RuntimeError("The 'openpyxl' library is required to read .xlsx files.\nInstall it with: pip install openpyxl")
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]

    def _analyze_import(self, all_rows, attrs, rels, lookthrough_names):
        """Classify each column and decide whether the file can be imported."""
        attr_by_safe = {a['safe']: a for a in attrs}
        rel_by_safe = {r['safe']: r for r in rels}
        look_safe = {sanitize_name(n) for n in lookthrough_names}

        headers_raw = all_rows[0] if all_rows else []
        col_map = []      # one entry per column: (kind, obj)
        columns = []      # (display, label, kind) for the preview
        seen = set()
        duplicates = []
        id_col_index = None
        mapped_attr_safes = set()

        for idx, h in enumerate(headers_raw):
            disp = "" if h is None else str(h).strip()
            s = sanitize_name(disp) if disp else ""
            if not disp:
                col_map.append(('blank', None)); continue
            if s == "id":
                id_col_index = idx
                col_map.append(('id', None))
                columns.append((disp, "ID (row key for update)", "key")); continue
            if s in attr_by_safe:
                if s in seen: duplicates.append(disp)
                seen.add(s); mapped_attr_safes.add(s)
                a = attr_by_safe[s]
                col_map.append(('attr', a))
                label = "File (loaded from this sheet's folder)" if a['type'] == 'file' else "Attribute"
                columns.append((disp, label, "ok")); continue
            if s in rel_by_safe:
                if s in seen: duplicates.append(disp)
                seen.add(s)
                col_map.append(('rel', rel_by_safe[s]))
                columns.append((disp, "Relationship", "ok")); continue
            if s in look_safe:
                col_map.append(('look', None))
                columns.append((disp, "Look-through (computed, ignored)", "ignore")); continue
            col_map.append(('unknown', None))
            columns.append((disp, "Unrecognized (ignored)", "ignore"))

        data_count = sum(1 for r in all_rows[1:] if any(c is not None and str(c).strip() != "" for c in r))
        missing_required = [a['name'] for a in attrs if a['required'] and a['safe'] not in mapped_attr_safes]
        has_recognized = any(k in ('attr', 'rel') for k, _ in col_map)

        problems = []
        if duplicates:
            problems.append("Duplicate columns map to the same field: " + ", ".join(duplicates))
        if missing_required:
            problems.append("Required columns are missing: " + ", ".join(missing_required))
        if data_count == 0:
            problems.append("No data rows found.")
        if not has_recognized:
            problems.append("No recognized attribute or relationship columns.")

        return {
            'columns': columns, 'col_map': col_map, 'id_col_index': id_col_index,
            'data_count': data_count, 'missing_required': missing_required,
            'duplicates': duplicates, 'problems': problems, 'can_import': not problems,
        }

    def import_data(self):
        if not self.current_view_name: return

        file_path, _ = QFileDialog.getOpenFileName(
            self, "Import Data", "",
            "Data Files (*.xlsx *.csv);;Excel Files (*.xlsx);;CSV Files (*.csv)")
        if not file_path: return

        try:
            all_rows = self._read_table_file(file_path)
        except Exception as e:
            QMessageBox.critical(self, "Import Failed", f"Could not read the file:\n\n{e}")
            return

        if not all_rows:
            QMessageBox.warning(self, "Import", "The file appears to be empty.")
            return

        with db_session(self.db_path) as conn:
            attrs, rels, lookthrough_names = self._get_import_schema(conn.cursor())

        analysis = self._analyze_import(all_rows, attrs, rels, lookthrough_names)
        analysis['file'] = os.path.basename(file_path)

        dlg = ImportPreviewDialog(self, analysis)
        if not dlg.exec(): return

        self._run_import(all_rows, attrs, rels, analysis, os.path.dirname(os.path.abspath(file_path)))

    def _run_import(self, all_rows, attrs, rels, analysis, source_dir):
        col_map = analysis['col_map']
        id_idx = analysis['id_col_index']
        table = self.current_table_name

        copied_abs = []     # files copied into storage this run (deleted on failure)
        old_to_trash = []   # replaced file references to retire after a successful commit
        try:
            with db_session(self.db_path) as conn:
                conn.execute("PRAGMA foreign_keys = 1")
                cur = conn.cursor()

                # --- snapshots for validation ---
                cur.execute(f"SELECT id FROM {qid(table)}")
                existing_ids = {r[0] for r in cur.fetchall()}

                uniq_maps = {}  # safe_col -> {value: owner_id}
                for a in attrs:
                    if a['unique']:
                        cur.execute(f"SELECT {qid(a['safe'])}, id FROM {qid(table)} WHERE {qid(a['safe'])} IS NOT NULL")
                        uniq_maps[a['safe']] = {v: i for v, i in cur.fetchall()}

                rel_data = {}  # safe -> {'ids': set, 'title_map': {lower: [ids]} or None, 'junc': str, 'name': str}
                for r in rels:
                    cur.execute(f"SELECT id FROM {qid(r['target_table'])}")
                    tids = {x[0] for x in cur.fetchall()}
                    title_map = None
                    if len(r['title_cols']) == 1:
                        title_map = {}
                        try:
                            cur.execute(f"SELECT [ID], {qid(r['title_cols'][0])} FROM {qid(r['base_view'])}")
                            for tid, tval in cur.fetchall():
                                if tval is not None and str(tval).strip() != "":
                                    title_map.setdefault(str(tval).strip().lower(), []).append(tid)
                        except sqlite3.OperationalError:
                            title_map = None
                    rel_data[r['safe']] = {'ids': tids, 'title_map': title_map,
                                           'junc': r['junc_table'], 'name': r['name'],
                                           'rel_type': r['rel_type']}

                # --- validate every row first (nothing written yet) ---
                file_unique_used = {}  # (safe, value) -> row_number
                seen_ids = set()
                ops = []

                for rn, row in enumerate(all_rows[1:], start=2):
                    if not any(c is not None and str(c).strip() != "" for c in row):
                        continue

                    mode, target_id = 'insert', None
                    if id_idx is not None and id_idx < len(row):
                        idraw = row[id_idx]
                        if idraw is not None and str(idraw).strip() != "":
                            try:
                                target_id = int(float(str(idraw).strip()))
                            except (ValueError, TypeError):
                                raise ValueError(f"Row {rn}: invalid ID '{idraw}'.")
                            if target_id in seen_ids:
                                raise ValueError(f"Row {rn}: ID {target_id} appears more than once in the file.")
                            seen_ids.add(target_id)
                            if target_id in existing_ids:
                                mode = 'update'
                            else:
                                mode, target_id = 'insert', None  # unknown ID -> add as new

                    set_cols, set_vals, rels_in_row, file_ops_row = [], [], {}, []

                    for idx, (kind, obj) in enumerate(col_map):
                        if idx >= len(row):
                            break
                        if kind == 'attr' and obj['type'] == 'file':
                            a = obj
                            val = self._cell_to_value(row[idx])
                            if val is None:
                                # Empty cell = keep the existing attachment (do not clear).
                                if a['required'] and mode == 'insert':
                                    raise ValueError(f"Row {rn}: '{a['name']}' is required but empty.")
                                continue
                            fname = os.path.basename(val)
                            src = os.path.join(source_dir, fname)
                            if not os.path.isfile(src):
                                raise ValueError(f"Row {rn}: file '{fname}' for '{a['name']}' was not found next to the import file.")
                            new_stored = make_stored_filename(fname)
                            set_cols.append(a['safe'])
                            set_vals.append(new_stored)
                            file_ops_row.append((a['safe'], src, new_stored))
                            continue
                        if kind == 'attr':
                            a = obj
                            val = self._cell_to_value(row[idx])
                            if val is None:
                                if a['required']:
                                    raise ValueError(f"Row {rn}: '{a['name']}' is required but empty.")
                                val_final = None
                            else:
                                ok, val_final = safe_convert(val, a['type'])
                                if not ok:
                                    raise ValueError(f"Row {rn}: cannot convert '{val}' for '{a['name']}' ({a['type']}).")
                                if a['type'] == 'discrete' and a['options'] is not None and val_final not in a['options']:
                                    raise ValueError(f"Row {rn}: '{val_final}' is not a valid option for '{a['name']}'.")
                                if a['type'] == 'matrix' and a['matrix_count']:
                                    try:
                                        parsed = ast.literal_eval(val_final)
                                    except (ValueError, SyntaxError):
                                        parsed = None
                                    if not isinstance(parsed, list) or len(parsed) != a['matrix_count']:
                                        raise ValueError(f"Row {rn}: '{a['name']}' must be a list of exactly {a['matrix_count']} column list(s).")
                                if a['unique'] and val_final is not None:
                                    owner = uniq_maps.get(a['safe'], {}).get(val_final)
                                    if owner is not None and owner != target_id:
                                        raise ValueError(f"Row {rn}: '{a['name']}' value '{val_final}' already exists in the database.")
                                    key = (a['safe'], val_final)
                                    if key in file_unique_used:
                                        raise ValueError(f"Row {rn}: '{a['name']}' value '{val_final}' is duplicated in the file (also row {file_unique_used[key]}).")
                                    file_unique_used[key] = rn
                            set_cols.append(a['safe'])
                            set_vals.append(val_final)

                        elif kind == 'rel':
                            r = obj
                            val = self._cell_to_value(row[idx])
                            rd = rel_data[r['safe']]
                            tids = []
                            if val is not None:
                                for tok in val.split(','):
                                    tok = tok.strip()
                                    if tok == "":
                                        continue
                                    if re.match(r'^-?\d+$', tok):
                                        tid = int(tok)
                                        if tid not in rd['ids']:
                                            raise ValueError(f"Row {rn}: '{r['name']}' target ID {tid} does not exist.")
                                    elif rd['title_map'] is not None:
                                        matches = rd['title_map'].get(tok.lower())
                                        if not matches:
                                            raise ValueError(f"Row {rn}: '{r['name']}' has no object titled '{tok}'.")
                                        if len(matches) > 1:
                                            raise ValueError(f"Row {rn}: '{r['name']}' title '{tok}' is ambiguous ({len(matches)} matches); use a numeric ID.")
                                        tid = matches[0]
                                    else:
                                        raise ValueError(f"Row {rn}: '{r['name']}' must use numeric IDs (target has no single title column).")
                                    tids.append(tid)
                            # De-duplicate (preserving order) so '1,1' can't create twin links,
                            # and enforce one_to_many cardinality that the object editor also enforces.
                            unique_tids = list(dict.fromkeys(tids))
                            if rd['rel_type'] == 'one_to_many' and len(unique_tids) > 1:
                                raise ValueError(
                                    f"Row {rn}: '{r['name']}' is a one-to-many relationship and accepts "
                                    f"a single target, but {len(unique_tids)} were given.")
                            rels_in_row[r['safe']] = unique_tids  # column present -> replace links

                    # Required relationships: a new object must supply a link; an update
                    # may omit the column (keep existing) but must not clear it.
                    for r in rels:
                        if not r.get('required'):
                            continue
                        present = r['safe'] in rels_in_row
                        has_link = present and len(rels_in_row[r['safe']]) > 0
                        if mode == 'insert' and not has_link:
                            raise ValueError(f"Row {rn}: '{r['name']}' is a required relationship — provide at least one target.")
                        if mode == 'update' and present and not has_link:
                            raise ValueError(f"Row {rn}: '{r['name']}' is a required relationship and cannot be cleared.")

                    ops.append({'mode': mode, 'id': target_id, 'cols': set_cols, 'vals': set_vals,
                                'rels': rels_in_row, 'files': file_ops_row})

                # --- apply atomically ---
                inserted = updated = 0
                files_dir = files_dir_for(self.db_path, create=True) if any(op['files'] for op in ops) else None
                conn.execute("BEGIN IMMEDIATE")
                for op in ops:
                    # On update, remember the file(s) we're about to replace so we can trash them.
                    if op['mode'] == 'update' and op['files']:
                        cols = [f[0] for f in op['files']]
                        sel = ", ".join(qid(c) for c in cols)
                        oldrow = cur.execute(f"SELECT {sel} FROM {qid(table)} WHERE id = ?", (op['id'],)).fetchone()
                        if oldrow:
                            old_to_trash.extend([v for v in oldrow if v])

                    # Copy new attachments into storage before the row write.
                    for (safe_col, src, new_stored) in op['files']:
                        dest_abs = os.path.join(files_dir, new_stored)
                        shutil.copy2(src, dest_abs)
                        copied_abs.append(dest_abs)

                    if op['mode'] == 'update':
                        oid = op['id']
                        if op['cols']:
                            set_clause = ", ".join(f"{qid(c)} = ?" for c in op['cols'])
                            cur.execute(f"UPDATE {qid(table)} SET {set_clause} WHERE id = ?", op['vals'] + [oid])
                        updated += 1
                    else:
                        if op['cols']:
                            ph = ", ".join(["?"] * len(op['cols']))
                            cur.execute(f"INSERT INTO {qid(table)} ({', '.join(qid(c) for c in op['cols'])}) VALUES ({ph})", op['vals'])
                        else:
                            cur.execute(f"INSERT INTO {qid(table)} DEFAULT VALUES")
                        oid = cur.lastrowid
                        inserted += 1

                    for rsafe, tids in op['rels'].items():
                        junc = rel_data[rsafe]['junc']
                        cur.execute(f"DELETE FROM {qid(junc)} WHERE source_id = ?", (oid,))
                        for tid in tids:
                            cur.execute(f"INSERT INTO {qid(junc)} (source_id, target_id) VALUES (?, ?)", (oid, tid))

                conn.commit()

            for old in old_to_trash:
                trash_stored_file(self.db_path, old)

            self.build_and_exec_query()
            QMessageBox.information(self, "Import Complete",
                f"Added {inserted} and updated {updated} object(s).")
        except Exception as e:
            for path in copied_abs:
                try:
                    os.remove(path)
                except OSError:
                    pass
            QMessageBox.critical(self, "Import Failed",
                f"Import aborted. No changes were made to the database.\n\nReason:\n{e}")
